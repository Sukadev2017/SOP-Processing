from __future__ import annotations

import hashlib
import re
import uuid

from pathlib import Path

from docx import Document

from .models import (
    CanonicalDocument,
    Element,
    Section,
    TextRun,
)


def sha256(
    path: Path,
):

    digest = hashlib.sha256()

    with path.open("rb") as file:

        for chunk in iter(
            lambda: file.read(
                1024 * 1024
            ),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest()


class DOCXExtractor:

    def __init__(
        self,
        template_config=None,
    ):

        config = (
            template_config
            or {}
        )

        heading_config = (
            config.get(
                "heading",
                {}
            )
        )

        self.max_level = (
            heading_config.get(
                "max_level",
                6,
            )
        )

        pattern = (
            heading_config.get(
                "numbered_heading_regex",
                (
                    r"^\s*"
                    r"(\d+(?:\.\d+)*)"
                    r"[\s\.\-]+"
                    r"(.+)$"
                ),
            )
        )

        self.numbered_heading = (
            re.compile(pattern)
        )

        self.uppercase_max = (
            heading_config.get(
                "uppercase_short_heading_max_chars",
                120,
            )
        )

    def _heading_info(
        self,
        paragraph,
    ):

        text = (
            paragraph.text.strip()
        )

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        # ---------------------------------------------
        # Best signal:
        # actual Word Heading style
        # ---------------------------------------------

        style_match = re.match(
            r"Heading\s+(\d+)",
            style_name,
            re.IGNORECASE,
        )

        if style_match:

            level = min(
                int(
                    style_match.group(1)
                ),
                self.max_level,
            )

            number_match = (
                self.numbered_heading
                .match(text)
            )

            if number_match:

                number = (
                    number_match.group(1)
                )

                heading = (
                    number_match
                    .group(2)
                    .strip()
                )

            else:

                number = None
                heading = text

            return (
                level,
                number,
                heading,
            )

        # ---------------------------------------------
        # Numbered headings
        # ---------------------------------------------

        match = (
            self.numbered_heading
            .match(text)
        )

        if match:

            number = (
                match.group(1)
            )

            heading = (
                match.group(2)
                .strip()
            )

            level = min(
                number.count(".")
                + 1,
                self.max_level,
            )

            return (
                level,
                number,
                heading,
            )

        # ---------------------------------------------
        # Fallback for controlled templates
        # containing uppercase headings.
        # ---------------------------------------------

        letters = re.sub(
            r"[^A-Za-z]+",
            "",
            text,
        )

        if (
            text
            and letters
            and len(text)
            <= self.uppercase_max
            and text
            == text.upper()
        ):

            return (
                1,
                None,
                text,
            )

        return None

    def extract(
        self,
        path: Path,
        document_type=None,
    ):

        document = Document(
            str(path)
        )

        elements = []

        sections = []

        order = 0

        current_section = None

        # =============================================
        # PARAGRAPHS
        # =============================================

        for paragraph in (
            document.paragraphs
        ):

            text = (
                paragraph.text.strip()
            )

            if not text:
                continue

            order += 1

            element_id = (
                f"DOCX-TXT-"
                f"{order:06d}"
            )

            runs = []

            for run in (
                paragraph.runs
            ):

                color = None

                if (
                    run.font.color
                    and run.font.color.rgb
                ):

                    color = (
                        f"#{run.font.color.rgb}"
                    )

                runs.append(
                    TextRun(
                        text=run.text,

                        font=(
                            run.font.name
                        ),

                        size=(
                            run.font.size.pt
                            if run.font.size
                            else None
                        ),

                        color=color,

                        bold=bool(
                            run.bold
                        ),

                        italic=bool(
                            run.italic
                        ),
                    )
                )

            heading_info = (
                self._heading_info(
                    paragraph
                )
            )

            metadata = {

                "style": (
                    paragraph.style.name
                    if paragraph.style
                    else None
                )
            }

            element_type = "text"

            if heading_info:

                (
                    level,
                    number,
                    heading,
                ) = heading_info

                element_type = (
                    "heading"
                )

                metadata.update(
                    {
                        "is_heading":
                            True,

                        "heading_level":
                            level,

                        "heading_number":
                            number,
                    }
                )

                current_section = (
                    Section(
                        section_id=(
                            f"TSEC-"
                            f"{len(sections)+1:03d}"
                        ),

                        heading=heading,

                        number=number,

                        level=level,

                        order=(
                            len(sections)
                            + 1
                        ),

                        heading_element_id=(
                            element_id
                        ),
                    )
                )

                sections.append(
                    current_section
                )

            element = Element(

                element_id=element_id,

                type=element_type,

                page=0,

                order=order,

                text=text,

                runs=runs,

                metadata=metadata,

                provenance={
                    "extractor":
                        "python-docx"
                },
            )

            elements.append(
                element
            )

            if (
                current_section
                and not heading_info
            ):

                current_section.element_ids.append(
                    element_id
                )

        # =============================================
        # TABLES
        # =============================================

        for (
            table_index,
            table,
        ) in enumerate(
            document.tables,
            1,
        ):

            order += 1

            element_id = (
                f"DOCX-TABLE-"
                f"{table_index:03d}"
            )

            rows = [

                [
                    cell.text.strip()

                    for cell
                    in row.cells
                ]

                for row
                in table.rows
            ]

            elements.append(
                Element(
                    element_id=element_id,

                    type="table",

                    page=0,

                    order=order,

                    text="\n".join(
                        " | ".join(row)
                        for row
                        in rows
                    ),

                    metadata={
                        "rows":
                            rows
                    },

                    provenance={
                        "extractor":
                            "python-docx"
                    },
                )
            )

        # =============================================
        # BUILD SECTION TEXT
        # =============================================

        element_lookup = {

            element.element_id:
                element

            for element
            in elements
        }

        for section in sections:

            section.text = "\n".join(

                element_lookup[
                    element_id
                ].text

                for element_id
                in section.element_ids

                if element_id
                in element_lookup
            )

        return CanonicalDocument(

            document_id=(
                "DOC-"
                + uuid.uuid4()
                .hex[:12]
                .upper()
            ),

            source_file=(
                path.name
            ),

            source_sha256=(
                sha256(path)
            ),

            document_type=(
                document_type
            ),

            elements=elements,

            sections=sections,
        )
