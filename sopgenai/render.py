from __future__ import annotations

import shutil

from pathlib import Path

from docx import Document

from docx.oxml import (
    OxmlElement,
)

from docx.text.paragraph import (
    Paragraph,
)


class ControlledDOCXRenderer:

    def render(
        self,
        template_path,
        template_document,
        mappings,
        output_path,
    ):

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        shutil.copy2(
            template_path,
            output_path,
        )

        document = Document(
            str(output_path)
        )

        mapping_lookup = {

            mapping.target_id:
                mapping

            for mapping
            in mappings
        }

        # ---------------------------------------------
        # Match target sections using the actual
        # extracted template headings.
        # ---------------------------------------------

        target_sections = {

            section.heading_element_id:
                section

            for section
            in template_document.sections

            if (
                section.heading_element_id
            )
        }

        paragraph_number = 0

        for paragraph in (
            document.paragraphs
        ):

            if not (
                paragraph.text.strip()
            ):

                continue

            paragraph_number += 1

            element_id = (
                f"DOCX-TXT-"
                f"{paragraph_number:06d}"
            )

            section = (
                target_sections.get(
                    element_id
                )
            )

            if not section:

                continue

            target_id = (
                section.number
                or section.section_id
            )

            mapping = (
                mapping_lookup.get(
                    target_id
                )
            )

            if not mapping:

                continue

            content = (
                mapping
                .transformed_content
                .strip()
            )

            if not content:

                continue

            new_xml = (
                OxmlElement(
                    "w:p"
                )
            )

            paragraph._p.addnext(
                new_xml
            )

            new_paragraph = (
                Paragraph(
                    new_xml,
                    paragraph._parent,
                )
            )

            try:

                new_paragraph.style = (
                    document.styles[
                        "Normal"
                    ]
                )

            except KeyError:

                pass

            new_paragraph.add_run(
                content
            )

        document.save(
            str(output_path)
        )

        return output_path
